from __future__ import annotations

import ast
import json
import os
import shutil
import tempfile
import unittest
import unicodedata
import uuid
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


APP_PATH = Path(__file__).resolve().parents[1] / "app.py"


def _load_parties_helpers():
    tree = ast.parse(APP_PATH.read_text(encoding="utf-8-sig"))
    wanted = {
        "safe_text",
        "party_code",
        "party_folder_name",
        "unique_folder_path",
        "_same_physical_path",
        "_party_family_prefix",
        "_party_family_dirs",
        "_party_code_from_folder_rel",
        "_is_party_suffix_variant",
        "_safe_child_path",
        "rename_party_dir",
        "load_parties",
        "save_parties",
        "export_parties_xlsx",
        "ensure_party_dirs_on_roots",
        "write_parties_log",
        "normalize_and_validate_parties",
        "apply_parties_update",
    }
    module = ast.Module(
        body=[node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name in wanted],
        type_ignores=[],
    )
    ast.fix_missing_locations(module)

    def pj(path: str, *parts) -> str:
        base = str(path).replace("/", "\\").rstrip("\\").strip()
        clean_parts = [str(p).strip("\\/ ") for p in parts if p]
        return "\\".join([base] + clean_parts) if clean_parts else base

    def load_json(path, default=None):
        try:
            return json.loads(Path(path).read_text(encoding="utf-8"))
        except Exception:
            return default if default is not None else {}

    def save_json(path: str, data):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        Path(path).write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")

    def sanitize_filename(name: str, max_len: int = 120) -> str:
        name = unicodedata.normalize("NFKD", str(name or "Piece"))
        name = name.encode("ascii", "ignore").decode("ascii")
        return name[:max_len].strip(" ._") or "Piece"

    ns = {
        "Path": Path,
        "Workbook": Workbook,
        "get_column_letter": get_column_letter,
        "datetime": datetime,
        "json": json,
        "os": os,
        "re": __import__("re"),
        "pj": pj,
        "load_json": load_json,
        "save_json": save_json,
        "sanitize_filename": sanitize_filename,
        "effective_nas_affaire_root": lambda cfg, aff_id: ((cfg or {}).get("roots") or {}).get("nas") or "",
        "pcfixe_unc_root_for_laptop": lambda cfg, aff_id: ((cfg or {}).get("roots") or {}).get("pcfixe") or "",
    }
    exec(compile(module, str(APP_PATH), "exec"), ns)
    return ns


class PartiesManagementTests(unittest.TestCase):
    def setUp(self):
        temp_root = Path(os.environ.get("LLM_ASSISTANT_TEST_TMP") or r"C:\CodexWorkspace\.tmp_llm_assistant_tests")
        temp_root.mkdir(exist_ok=True)
        self.root = temp_root / f"parties_{uuid.uuid4().hex}"
        self.root.mkdir()
        self.cfg_dir = self.root / "_Config"
        self.cfg_dir.mkdir()
        (self.root / "AB_Organisation_expertise").mkdir()
        (self.root / "AA_Expert_Admin" / "_Logs").mkdir(parents=True)
        self.ns = _load_parties_helpers()
        self.project_config = {
            "roots": {
                "laptop": str(self.root),
                "nas": str(self.root),
                "pcfixe": str(self.root),
            }
        }

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def _write_parties(self, parties):
        (self.cfg_dir / "parties.json").write_text(
            json.dumps({"parties": parties}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    def _read_parties(self):
        return json.loads((self.cfg_dir / "parties.json").read_text(encoding="utf-8"))["parties"]

    def _party(self, code, name=None, folder_rel=None):
        name = name or f"Partie {code:02d}"
        folder_rel = folder_rel or f"{code:02d}_Partie_{code:02d}_{name}"
        return {
            "code_partie": code,
            "nom": name,
            "representant": "",
            "avocat": "",
            "notes": "",
            "folder_rel": folder_rel,
            "history": [],
        }

    def test_creation_and_rename_still_work(self):
        res_create = self.ns["apply_parties_update"](
            str(self.root),
            str(self.cfg_dir),
            [{"code_partie": 1, "nom": "Alpha", "representant": "", "avocat": "", "notes": ""}],
            aff_id="2099-J01",
            titre="Test",
            project_config=self.project_config,
            export_xlsx=False,
        )
        self.assertEqual([1], res_create["created_codes"])
        old_folder = self._read_parties()[0]["folder_rel"]
        self.assertTrue((self.root / old_folder).is_dir())

        res_rename = self.ns["apply_parties_update"](
            str(self.root),
            str(self.cfg_dir),
            [{"code_partie": 1, "nom": "Beta", "representant": "", "avocat": "", "notes": ""}],
            aff_id="2099-J01",
            titre="Test",
            project_config=self.project_config,
            export_xlsx=False,
        )
        new_folder = self._read_parties()[0]["folder_rel"]
        self.assertTrue(res_rename["renamed"])
        self.assertFalse((self.root / old_folder).exists())
        self.assertTrue((self.root / new_folder).is_dir())

    def test_apply_without_modification_is_idempotent_and_creates_no_suffix(self):
        rows = [{"code_partie": 1, "nom": "Alpha", "representant": "", "avocat": "", "notes": ""}]

        self.ns["apply_parties_update"](
            str(self.root),
            str(self.cfg_dir),
            rows,
            aff_id="2099-J01",
            titre="Test",
            project_config=self.project_config,
            export_xlsx=False,
        )
        for _ in range(3):
            res = self.ns["apply_parties_update"](
                str(self.root),
                str(self.cfg_dir),
                rows,
                aff_id="2099-J01",
                titre="Test",
                project_config=self.project_config,
                export_xlsx=False,
            )
            self.assertFalse(res["renamed"])

        family = sorted(p.name for p in self.root.iterdir() if p.is_dir() and p.name.startswith("01_Partie_01_"))
        self.assertEqual(["01_Partie_01_Alpha"], family)
        self.assertFalse(any("__" in name for name in family))

    def test_metadata_update_does_not_create_or_rename_folder(self):
        party = self._party(1, "Alpha")
        self._write_parties([party])
        (self.root / party["folder_rel"]).mkdir()

        res = self.ns["apply_parties_update"](
            str(self.root),
            str(self.cfg_dir),
            [{"code_partie": 1, "nom": "Alpha", "representant": "Rep", "avocat": "Avocat", "notes": "Note"}],
            aff_id="2099-J01",
            titre="Test",
            project_config=self.project_config,
            export_xlsx=False,
        )

        self.assertFalse(res["renamed"])
        self.assertTrue((self.root / party["folder_rel"]).is_dir())
        self.assertEqual(party["folder_rel"], self._read_parties()[0]["folder_rel"])

    def test_rename_preserves_existing_files(self):
        party = self._party(1, "Alpha")
        self._write_parties([party])
        folder = self.root / party["folder_rel"]
        (folder / "Sous").mkdir(parents=True)
        (folder / "Sous" / "piece.pdf").write_text("contenu", encoding="utf-8")

        self.ns["apply_parties_update"](
            str(self.root),
            str(self.cfg_dir),
            [{"code_partie": 1, "nom": "Beta", "representant": "", "avocat": "", "notes": ""}],
            aff_id="2099-J01",
            titre="Test",
            project_config=self.project_config,
            export_xlsx=False,
        )

        new_folder = self.root / "01_Partie_01_Beta"
        self.assertFalse(folder.exists())
        self.assertEqual("contenu", (new_folder / "Sous" / "piece.pdf").read_text(encoding="utf-8"))

    def test_destination_occupied_by_other_folder_blocks_without_suffix(self):
        party = self._party(1, "Alpha")
        self._write_parties([party])
        (self.root / party["folder_rel"]).mkdir()
        occupied = self.root / "01_Partie_01_Beta"
        occupied.mkdir()

        before = (self.cfg_dir / "parties.json").read_text(encoding="utf-8")
        with self.assertRaises(ValueError) as ctx:
            self.ns["apply_parties_update"](
                str(self.root),
                str(self.cfg_dir),
                [{"code_partie": 1, "nom": "Beta", "representant": "", "avocat": "", "notes": ""}],
                aff_id="2099-J01",
                titre="Test",
                project_config=self.project_config,
                export_xlsx=True,
            )

        self.assertIn("destination canonique deja occupee", str(ctx.exception))
        self.assertEqual(before, (self.cfg_dir / "parties.json").read_text(encoding="utf-8"))
        self.assertTrue((self.root / party["folder_rel"]).is_dir())
        self.assertTrue(occupied.is_dir())
        self.assertFalse(any(p.name.endswith("__2") for p in self.root.iterdir() if p.is_dir()))

    def test_creation_blocks_when_same_code_family_already_exists(self):
        (self.root / "01_Partie_01_Ancien").mkdir()

        with self.assertRaises(ValueError) as ctx:
            self.ns["apply_parties_update"](
                str(self.root),
                str(self.cfg_dir),
                [{"code_partie": 1, "nom": "Alpha", "representant": "", "avocat": "", "notes": ""}],
                aff_id="2099-J01",
                titre="Test",
                project_config=self.project_config,
                export_xlsx=False,
            )

        self.assertIn("dossier de la meme famille", str(ctx.exception))
        self.assertFalse((self.root / "01_Partie_01_Alpha").exists())
        self.assertFalse(any(p.name.endswith("__2") for p in self.root.iterdir() if p.is_dir()))

    def test_ensure_party_dirs_blocks_family_conflict_on_other_roots(self):
        party = self._party(1, "Alpha")
        other_root = self.root / "other_root"
        other_root.mkdir()
        (other_root / "01_Partie_01_Old").mkdir()
        cfg = {"roots": {"nas": str(other_root), "pcfixe": str(self.root)}}

        result = self.ns["ensure_party_dirs_on_roots"](str(self.root), cfg, "2099-J01", [party])

        self.assertTrue(any(item.get("target") == "nas" and "autre dossier" in item.get("error", "") for item in result["errors"]))
        self.assertFalse((other_root / party["folder_rel"]).exists())
        self.assertFalse(any(p.name.endswith("__2") for p in other_root.iterdir() if p.is_dir()))

    def test_excel_and_parties_json_updated_without_suffix_after_apply(self):
        rows = [{"code_partie": 1, "nom": "Alpha", "representant": "Rep", "avocat": "Avocat", "notes": "Note"}]

        res = self.ns["apply_parties_update"](
            str(self.root),
            str(self.cfg_dir),
            rows,
            aff_id="2099-J01",
            titre="Test",
            project_config=self.project_config,
            export_xlsx=True,
        )

        party = self._read_parties()[0]
        self.assertEqual("01_Partie_01_Alpha", party["folder_rel"])
        self.assertNotIn("__", party["folder_rel"])
        self.assertTrue(Path(res["xlsx_path"]).is_file())
        wb = load_workbook(res["xlsx_path"], read_only=True, data_only=True)
        try:
            rows_out = list(wb.active.iter_rows(min_row=5, values_only=True))
        finally:
            wb.close()
        self.assertEqual("01", rows_out[0][0])
        self.assertEqual("01_Partie_01_Alpha", rows_out[0][5])


if __name__ == "__main__":
    unittest.main()
