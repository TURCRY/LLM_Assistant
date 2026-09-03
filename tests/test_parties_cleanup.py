from __future__ import annotations

import ast
import hashlib
import json
import os
import shutil
import sqlite3
import tempfile
import unittest
import unicodedata
import uuid
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


APP_PATH = Path(__file__).resolve().parents[1] / "app.py"


def _load_cleanup_helpers():
    tree = ast.parse(APP_PATH.read_text(encoding="utf-8-sig"))
    wanted = {
        "safe_text",
        "compact_spaces",
        "party_code",
        "party_folder_name",
        "unique_folder_path",
        "_same_physical_path",
        "_party_family_prefix",
        "_party_family_dirs",
        "_party_code_from_folder_rel",
        "_is_party_suffix_variant",
        "_safe_child_path",
        "rename_party_dir",
        "load_parties",
        "save_parties",
        "export_parties_xlsx",
        "ensure_party_dirs_on_roots",
        "party_roots_for_affaire",
        "_path_dedup_key",
        "_is_safe_party_folder_rel",
        "party_delete_selection_signature",
        "party_delete_confirmation_valid",
        "party_cleanup_selection_signature",
        "party_cleanup_confirmation_valid",
        "is_party_folder_candidate_name",
        "is_safe_party_cleanup_folder_rel",
        "dedupe_party_roots",
        "build_active_party_folder_index",
        "inspect_party_folder_path",
        "inspect_party_folder_tree",
        "sha256_file",
        "detect_party_folder_file_conflicts",
        "valid_expert_doc_id",
        "expert_doc_id_sort_key",
        "_documentary_code_from_expert_id",
        "_normalized_documentary_code",
        "_documentary_code_from_doc",
        "_documentary_identity_diagnostics",
        "_documentary_stable_party_key",
        "_documentary_filename",
        "_documentary_path_values",
        "deduce_code_source_from_paths",
        "normalize_source_code",
        "load_sqlite_documentary_rows",
        "build_documentary_consistency_diagnostic",
        "analyze_party_folder_duplicates",
        "load_party_folder_history",
        "analyze_residual_party_folders",
        "authoritative_party_mutation_root",
        "preflight_cleanup_residual_party_folders",
        "write_parties_cleanup_log",
        "cleanup_residual_party_folders",
        "preflight_delete_parties",
        "write_parties_delete_log",
        "delete_parties",
    }
    wanted_constants = {
        "PARTY_FOLDER_PATTERN",
        "PARTY_CLEANUP_ACTIVE",
        "PARTY_CLEANUP_RESIDUAL_EMPTY",
        "PARTY_CLEANUP_NON_EMPTY",
        "PARTY_CLEANUP_INCONSISTENT",
        "PARTY_REPAIR_OK",
        "PARTY_REPAIR_EMPTY_RESIDUES",
        "PARTY_REPAIR_TO_CONSOLIDATE",
        "PARTY_REPAIR_CONFLICT",
        "PARTY_REPAIR_INCONSISTENT",
    }
    body = []
    for node in tree.body:
        if isinstance(node, ast.FunctionDef) and node.name in wanted:
            body.append(node)
        elif isinstance(node, ast.Assign):
            names = {target.id for target in node.targets if isinstance(target, ast.Name)}
            if names & wanted_constants:
                body.append(node)
    module = ast.Module(body=body, type_ignores=[])
    ast.fix_missing_locations(module)

    def pj(path: str, *parts) -> str:
        base = str(path).replace("/", "\\").rstrip("\\").strip()
        clean_parts = [str(p).strip("\\/ ") for p in parts if p]
        return "\\".join([base] + clean_parts) if clean_parts else base

    def load_json(path, default=None):
        try:
            return json.loads(Path(path).read_text(encoding="utf-8"))
        except Exception:
            return default if default is not None else {}

    def save_json(path: str, data):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        Path(path).write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")

    def sanitize_filename(name: str, max_len: int = 120) -> str:
        name = unicodedata.normalize("NFKD", str(name or "Piece"))
        name = name.encode("ascii", "ignore").decode("ascii")
        return name[:max_len].strip(" ._") or "Piece"

    ns = {
        "Path": Path,
        "Workbook": Workbook,
        "get_column_letter": get_column_letter,
        "datetime": datetime,
        "hashlib": hashlib,
        "sqlite3": sqlite3,
        "json": json,
        "os": os,
        "re": __import__("re"),
        "pj": pj,
        "load_json": load_json,
        "save_json": save_json,
        "sanitize_filename": sanitize_filename,
        "effective_nas_affaire_root": lambda cfg, aff_id: ((cfg or {}).get("roots") or {}).get("nas") or "",
        "pcfixe_unc_root_for_laptop": lambda cfg, aff_id: ((cfg or {}).get("roots") or {}).get("pcfixe") or "",
        "affaire_sqlite_path_from_root": lambda root, cfg=None: Path(root) / "Documents.db",
    }
    exec(compile(module, str(APP_PATH), "exec"), ns)
    return ns


class PartiesCleanupTests(unittest.TestCase):
    def setUp(self):
        temp_root = Path(os.environ.get("LLM_ASSISTANT_TEST_TMP") or r"C:\CodexWorkspace\.tmp_llm_assistant_tests")
        temp_root.mkdir(exist_ok=True)
        self.base = temp_root / f"cleanup_{uuid.uuid4().hex}"
        self.laptop = self.base / "laptop"
        self.nas = self.base / "nas"
        self.pcfixe = self.base / "pcfixe"
        for root in (self.laptop, self.nas, self.pcfixe):
            (root / "_Config").mkdir(parents=True)
            (root / "AA_Expert_Admin" / "_Logs").mkdir(parents=True)
            (root / "AB_Organisation_expertise").mkdir(parents=True)
        self.cfg_dir = self.laptop / "_Config"
        self.ns = _load_cleanup_helpers()
        self.project_config = {
            "roots": {
                "laptop": str(self.laptop),
                "nas": str(self.nas),
                "pcfixe": str(self.pcfixe),
            }
        }

    def tearDown(self):
        shutil.rmtree(self.base, ignore_errors=True)

    def _party(self, code, name, folder_rel=None):
        return {
            "code_partie": code,
            "nom": name,
            "representant": "",
            "avocat": "",
            "notes": "",
            "folder_rel": folder_rel or f"{code:02d}_Partie_{code:02d}_{name}",
            "history": [],
        }

    def _write_parties(self, parties):
        (self.cfg_dir / "parties.json").write_text(
            json.dumps({"parties": parties}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    def _make_folder_all_roots(self, folder_rel):
        paths = []
        for root in (self.laptop, self.nas, self.pcfixe):
            path = root / folder_rel
            path.mkdir(parents=True)
            paths.append(path)
        return paths

    def _analyze(self, parties=None, cfg=None):
        return self.ns["analyze_residual_party_folders"](
            str(self.laptop),
            cfg or self.project_config,
            "2099-J01",
            parties if parties is not None else self.ns["load_parties"](str(self.cfg_dir)),
        )

    def _folder(self, analysis, folder_rel):
        for item in analysis["folders"]:
            if item["folder_rel"] == folder_rel:
                return item
        self.fail(f"Dossier non trouvé dans l'analyse: {folder_rel}")

    def test_active_folder_is_never_selectable(self):
        party = self._party(1, "Alpha")
        self._write_parties([party])
        self._make_folder_all_roots(party["folder_rel"])

        item = self._folder(self._analyze(), party["folder_rel"])

        self.assertEqual(self.ns["PARTY_CLEANUP_ACTIVE"], item["state"])
        self.assertFalse(item["selectable"])

    def test_empty_residual_on_all_roots_is_selectable(self):
        self._write_parties([])
        folder = "10_Partie_10_Residual"
        self._make_folder_all_roots(folder)

        item = self._folder(self._analyze(), folder)

        self.assertEqual(self.ns["PARTY_CLEANUP_RESIDUAL_EMPTY"], item["state"])
        self.assertTrue(item["selectable"])

    def test_file_on_one_root_blocks_cleanup(self):
        self._write_parties([])
        folder = "10_Partie_10_Residual"
        paths = self._make_folder_all_roots(folder)
        (paths[0] / "piece.pdf").write_text("x", encoding="utf-8")

        item = self._folder(self._analyze(), folder)

        self.assertIn(item["state"], {self.ns["PARTY_CLEANUP_NON_EMPTY"], self.ns["PARTY_CLEANUP_INCONSISTENT"]})
        self.assertFalse(item["selectable"])

    def test_subfolder_on_one_root_blocks_cleanup(self):
        self._write_parties([])
        folder = "10_Partie_10_Residual"
        paths = self._make_folder_all_roots(folder)
        (paths[1] / "SousDossier").mkdir()

        item = self._folder(self._analyze(), folder)

        self.assertIn(item["state"], {self.ns["PARTY_CLEANUP_NON_EMPTY"], self.ns["PARTY_CLEANUP_INCONSISTENT"]})
        self.assertFalse(item["selectable"])

    def test_missing_on_one_root_is_inconsistent(self):
        self._write_parties([])
        folder = "10_Partie_10_Residual"
        (self.laptop / folder).mkdir()
        (self.nas / folder).mkdir()

        item = self._folder(self._analyze(), folder)

        self.assertEqual(self.ns["PARTY_CLEANUP_INCONSISTENT"], item["state"])
        self.assertFalse(item["selectable"])

    def test_inaccessible_root_blocks_analysis_and_selection(self):
        self._write_parties([])
        folder = "10_Partie_10_Residual"
        (self.laptop / folder).mkdir()
        cfg = {"roots": {"nas": str(self.base / "missing"), "pcfixe": str(self.pcfixe)}}

        analysis = self._analyze(cfg=cfg)
        preflight = self.ns["preflight_cleanup_residual_party_folders"](
            str(self.laptop), cfg, "2099-J01", [], [folder]
        )

        self.assertFalse(analysis["ok"])
        self.assertEqual(self.ns["PARTY_CLEANUP_INCONSISTENT"], self._folder(analysis, folder)["state"])
        self.assertFalse(preflight["ok"])

    def test_path_outside_root_is_refused(self):
        preflight = self.ns["preflight_cleanup_residual_party_folders"](
            str(self.laptop), self.project_config, "2099-J01", [], [r"..\outside"]
        )

        self.assertFalse(preflight["ok"])
        self.assertTrue(any(".." in str(b.get("error")) for b in preflight["blockers"]))

    def test_bad_format_is_not_candidate(self):
        self._write_parties([])
        bad = "10_Partie_11_Bad"
        self._make_folder_all_roots(bad)

        analysis = self._analyze()

        self.assertNotIn(bad, {item["folder_rel"] for item in analysis["folders"]})

    def test_suffixes_have_no_special_cleanup_rule(self):
        self._write_parties([])
        for suffix in ("__2", "__4", "__27"):
            self._make_folder_all_roots(f"10_Partie_10_Residual{suffix}")

        analysis = self._analyze()

        states = {item["folder_rel"]: item["state"] for item in analysis["folders"]}
        self.assertEqual(self.ns["PARTY_CLEANUP_RESIDUAL_EMPTY"], states["10_Partie_10_Residual__2"])
        self.assertEqual(self.ns["PARTY_CLEANUP_RESIDUAL_EMPTY"], states["10_Partie_10_Residual__4"])
        self.assertEqual(self.ns["PARTY_CLEANUP_RESIDUAL_EMPTY"], states["10_Partie_10_Residual__27"])

    def test_equivalent_roots_are_deduplicated(self):
        self._write_parties([])
        folder = "10_Partie_10_Residual"
        (self.laptop / folder).mkdir()
        cfg = {"roots": {"nas": str(self.laptop), "pcfixe": str(self.laptop)}}

        analysis = self._analyze(cfg=cfg)
        preflight = self.ns["preflight_cleanup_residual_party_folders"](
            str(self.laptop), cfg, "2099-J01", [], [folder]
        )

        self.assertEqual(1, len(analysis["roots"]))
        self.assertTrue(preflight["ok"])
        self.assertEqual(1, len(preflight["targets"]))

    def test_multiple_valid_selection_deletes_selected_folders_on_nas_only(self):
        self._write_parties([])
        folders = ["10_Partie_10_A", "11_Partie_11_B"]
        for folder in folders:
            self._make_folder_all_roots(folder)

        res = self.ns["cleanup_residual_party_folders"](
            str(self.laptop), str(self.cfg_dir), folders, aff_id="2099-J01", project_config=self.project_config
        )

        self.assertEqual(2, len(res["removed_dirs"]))
        self.assertTrue(all(item["target"] == "nas" for item in res["removed_dirs"]))
        for folder in folders:
            self.assertTrue((self.laptop / folder).is_dir())
            self.assertFalse((self.nas / folder).exists())
            self.assertTrue((self.pcfixe / folder).is_dir())

    def test_one_blocked_target_in_multiple_selection_prevents_any_mutation(self):
        self._write_parties([])
        valid = "10_Partie_10_A"
        blocked = "11_Partie_11_B"
        self._make_folder_all_roots(valid)
        blocked_paths = self._make_folder_all_roots(blocked)
        (blocked_paths[2] / "piece.pdf").write_text("x", encoding="utf-8")

        with self.assertRaises(ValueError):
            self.ns["cleanup_residual_party_folders"](
                str(self.laptop), str(self.cfg_dir), [valid, blocked], aff_id="2099-J01", project_config=self.project_config
            )

        self.assertTrue(all((root / valid).is_dir() for root in (self.laptop, self.nas, self.pcfixe)))
        self.assertTrue(all((root / blocked).is_dir() for root in (self.laptop, self.nas, self.pcfixe)))

    def test_second_preflight_blocks_folder_that_became_non_empty(self):
        self._write_parties([])
        folder = "10_Partie_10_Residual"
        paths = self._make_folder_all_roots(folder)
        analysis = self._analyze()
        self.assertTrue(self._folder(analysis, folder)["selectable"])
        (paths[1] / "new.pdf").write_text("x", encoding="utf-8")

        with self.assertRaises(ValueError):
            self.ns["cleanup_residual_party_folders"](
                str(self.laptop), str(self.cfg_dir), [folder], aff_id="2099-J01", project_config=self.project_config
            )

        self.assertTrue(all((root / folder).is_dir() for root in (self.laptop, self.nas, self.pcfixe)))

    def test_second_preflight_blocks_folder_that_became_active(self):
        self._write_parties([])
        folder = "10_Partie_10_Residual"
        self._make_folder_all_roots(folder)
        analysis = self._analyze()
        self.assertTrue(self._folder(analysis, folder)["selectable"])
        self._write_parties([self._party(10, "Residual", folder)])

        with self.assertRaises(ValueError):
            self.ns["cleanup_residual_party_folders"](
                str(self.laptop), str(self.cfg_dir), [folder], aff_id="2099-J01", project_config=self.project_config
            )

        self.assertTrue(all((root / folder).is_dir() for root in (self.laptop, self.nas, self.pcfixe)))

    def test_cleanup_confirmation_signature_tracks_exact_folder_set(self):
        signature = self.ns["party_cleanup_selection_signature"]
        is_valid = self.ns["party_cleanup_confirmation_valid"]

        self.assertEqual(signature(["B", "A"]), signature(["A", "B"]))
        confirmed = signature(["A"])
        self.assertTrue(is_valid(["A"], True, confirmed))
        self.assertFalse(is_valid(["A", "B"], True, confirmed))
        self.assertFalse(is_valid([], True, confirmed))
        confirmed = signature(["A", "B"])
        self.assertTrue(is_valid(["B", "A"], True, confirmed))
        self.assertFalse(is_valid(["A"], True, confirmed))

    def test_cleanup_log_contains_essential_operation_details(self):
        self._write_parties([])
        folder = "10_Partie_10_Residual"
        self._make_folder_all_roots(folder)

        res = self.ns["cleanup_residual_party_folders"](
            str(self.laptop), str(self.cfg_dir), [folder], aff_id="2099-J01", project_config=self.project_config
        )
        payload = json.loads(Path(res["log_path"]).read_text(encoding="utf-8"))

        self.assertEqual("cleanup_residual_party_folders", payload["action"])
        self.assertEqual("2099-J01", payload["aff_id"])
        self.assertEqual([folder], payload["requested_selection"])
        self.assertTrue(payload["roots"])
        self.assertTrue(payload["preflight"]["ok"])
        self.assertTrue(payload["removed_dirs"])
        self.assertEqual(folder, payload["classification"][0]["folder_rel"])

    def test_parties_json_and_excel_are_unchanged_by_cleanup(self):
        parties = [self._party(1, "Active")]
        self._write_parties(parties)
        folder = "10_Partie_10_Residual"
        self._make_folder_all_roots(folder)
        xlsx_path = self.laptop / "AB_Organisation_expertise" / "Id_affaire_en_tete_dossier.xlsx"
        self.ns["export_parties_xlsx"](str(xlsx_path), "2099-J01", "Test", parties)
        json_before = (self.cfg_dir / "parties.json").read_bytes()
        xlsx_before = xlsx_path.read_bytes()

        self.ns["cleanup_residual_party_folders"](
            str(self.laptop), str(self.cfg_dir), [folder], aff_id="2099-J01", project_config=self.project_config
        )

        self.assertEqual(json_before, (self.cfg_dir / "parties.json").read_bytes())
        self.assertEqual(xlsx_before, xlsx_path.read_bytes())

    def test_history_log_is_reported_when_available(self):
        self._write_parties([])
        folder = "10_Partie_10_Old"
        self._make_folder_all_roots(folder)
        log_path = self.laptop / "AA_Expert_Admin" / "_Logs" / "parties_update_20990101_010101.json"
        log_path.write_text(
            json.dumps({"ts": "2099-01-01T01:01:01", "renamed": [{"code_partie": 10, "from": folder, "to": "10_Partie_10_New"}]}),
            encoding="utf-8",
        )

        item = self._folder(self._analyze(), folder)

        self.assertEqual(1, len(item["history"]))
        self.assertEqual("from", item["history"][0]["role"])

    def test_existing_active_delete_still_updates_registry(self):
        parties = [self._party(1, "Alpha"), self._party(2, "Beta")]
        self._write_parties(parties)
        for party in parties:
            (self.laptop / party["folder_rel"]).mkdir()
        cfg = {"roots": {"nas": str(self.laptop), "pcfixe": str(self.laptop)}}

        res = self.ns["delete_parties"](
            str(self.laptop),
            str(self.cfg_dir),
            [2],
            aff_id="2099-J01",
            titre="Test",
            project_config=cfg,
            export_xlsx=False,
        )
        remaining = json.loads((self.cfg_dir / "parties.json").read_text(encoding="utf-8"))["parties"]

        self.assertEqual([1], [p["code_partie"] for p in remaining])
        self.assertEqual([2], res["deleted_codes"])
        self.assertTrue(all(item["target"] == "nas" for item in res["removed_dirs"]))

    def _repair_analysis(self, parties=None, cfg=None):
        return self.ns["analyze_party_folder_duplicates"](
            str(self.laptop),
            cfg or self.project_config,
            "2099-J01",
            parties if parties is not None else self.ns["load_parties"](str(self.cfg_dir)),
        )

    def _repair_code(self, analysis, code):
        for item in analysis["codes"]:
            if item["code_partie"] == code:
                return item
        self.fail(f"Code non trouvé dans le diagnostic: {code}")

    def test_repair_diagnostic_canonical_only_is_ok(self):
        party = self._party(1, "Alpha")
        self._write_parties([party])
        self._make_folder_all_roots(party["folder_rel"])

        item = self._repair_code(self._repair_analysis(), 1)

        self.assertEqual(self.ns["PARTY_REPAIR_OK"], item["status"])
        self.assertEqual("01_Partie_01_Alpha", item["canonical_folder_rel"])
        self.assertFalse(item["repair_mutation_allowed"])

    def test_repair_diagnostic_groups_suffixes_by_code(self):
        party = self._party(17, "Lallou GUEZ")
        self._write_parties([party])
        self._make_folder_all_roots("17_Partie_17_Lallou GUEZ")
        self._make_folder_all_roots("17_Partie_17_Ancien Nom")
        self._make_folder_all_roots("17_Partie_17_Lallou GUEZ__2")

        item = self._repair_code(self._repair_analysis(), 17)

        self.assertEqual(self.ns["PARTY_REPAIR_EMPTY_RESIDUES"], item["status"])
        self.assertEqual(
            {"17_Partie_17_Ancien Nom", "17_Partie_17_Lallou GUEZ__2"},
            set(item["other_folder_rels"]),
        )

    def test_repair_diagnostic_suffix_with_data_requires_consolidation(self):
        party = self._party(1, "Alpha")
        self._write_parties([party])
        paths = self._make_folder_all_roots("01_Partie_01_Alpha__2")
        (paths[0] / "piece.pdf").write_text("data", encoding="utf-8")

        item = self._repair_code(self._repair_analysis(), 1)

        self.assertEqual(self.ns["PARTY_REPAIR_TO_CONSOLIDATE"], item["status"])
        self.assertEqual("01_Partie_01_Alpha", item["canonical_folder_rel"])

    def test_repair_diagnostic_detects_identical_duplicates(self):
        party = self._party(1, "Alpha")
        self._write_parties([party])
        canonical_paths = self._make_folder_all_roots("01_Partie_01_Alpha")
        suffix_paths = self._make_folder_all_roots("01_Partie_01_Alpha__2")
        (canonical_paths[0] / "piece.pdf").write_text("same", encoding="utf-8")
        (suffix_paths[0] / "piece.pdf").write_text("same", encoding="utf-8")

        item = self._repair_code(self._repair_analysis(), 1)

        self.assertEqual(self.ns["PARTY_REPAIR_TO_CONSOLIDATE"], item["status"])
        self.assertEqual(1, len(item["duplicates"]))
        self.assertFalse(item["conflicts"])

    def test_repair_diagnostic_detects_conflicting_same_relative_path(self):
        party = self._party(1, "Alpha")
        self._write_parties([party])
        canonical_paths = self._make_folder_all_roots("01_Partie_01_Alpha")
        suffix_paths = self._make_folder_all_roots("01_Partie_01_Alpha__2")
        (canonical_paths[0] / "Sous").mkdir()
        (suffix_paths[0] / "Sous").mkdir()
        (canonical_paths[0] / "Sous" / "piece.pdf").write_text("left", encoding="utf-8")
        (suffix_paths[0] / "Sous" / "piece.pdf").write_text("right", encoding="utf-8")

        item = self._repair_code(self._repair_analysis(), 1)

        self.assertEqual(self.ns["PARTY_REPAIR_CONFLICT"], item["status"])
        self.assertEqual("Sous/piece.pdf", item["conflicts"][0]["rel"])

    def test_repair_diagnostic_reports_documentary_code_path_divergence(self):
        party = self._party(1, "Alpha")
        self._write_parties([party])
        self._make_folder_all_roots(party["folder_rel"])
        conn = sqlite3.connect(str(self.laptop / "Documents.db"))
        conn.execute(
            """
            CREATE TABLE Documents (
                id_document INTEGER PRIMARY KEY,
                numero_expert TEXT NOT NULL,
                code_partie TEXT NOT NULL,
                nom_original TEXT,
                nom_cible TEXT,
                chemin_nas TEXT,
                chemin_local TEXT
            )
            """
        )
        conn.execute(
            """
            INSERT INTO Documents (numero_expert, code_partie, nom_original, nom_cible, chemin_nas, chemin_local)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                "01-001",
                "01",
                "piece.pdf",
                "piece.pdf",
                str(self.nas / "02_Partie_02_Beta" / "piece.pdf"),
                str(self.laptop / "01_Partie_01_Alpha" / "piece.pdf"),
            ),
        )
        conn.commit()
        conn.close()

        analysis = self._repair_analysis()
        item = self._repair_code(analysis, 1)

        self.assertEqual(1, item["documentaire"]["sqlite_documents_count"])
        self.assertEqual(["01-001"], item["documentaire"]["expert_doc_ids"])
        self.assertEqual("INCOHÉRENCE", item["documentaire"]["statut_documentaire"])
        self.assertEqual("02", item["documentaire"]["documents_chemin_autre_code"][0]["code_chemin"])

    def test_repair_diagnostic_inaccessible_root_blocks_mutation(self):
        party = self._party(1, "Alpha")
        self._write_parties([party])
        (self.laptop / party["folder_rel"]).mkdir()
        cfg = {"roots": {"nas": str(self.base / "missing"), "pcfixe": str(self.pcfixe)}}

        analysis = self._repair_analysis(cfg=cfg)
        item = self._repair_code(analysis, 1)

        self.assertFalse(analysis["ok"])
        self.assertEqual(self.ns["PARTY_REPAIR_INCONSISTENT"], item["status"])
        self.assertFalse(item["repair_mutation_allowed"])


if __name__ == "__main__":
    unittest.main()
